"""Download unchanged public CbC datasets locally and check ATO's catalogue."""
import argparse
import csv
import json
import re
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

import build_company_data as companies

CATALOGUE_URL = "https://data.gov.au/data/api/3/action/package_search"
ATO_ORGANIZATION = "australiantaxationoffice"
GUIDANCE_URL = "https://softwaredevelopers.ato.gov.au/PublicCBCreporting"
DOWNLOADS = [
    ("taxplorer-data.csv", "https://www.taxplorer.eu/taxplorer-data.csv", "EU Tax Observatory, Taxplorer"),
    ("taxplorer-data.xlsx", "https://www.taxplorer.eu/taxplorer-data.xlsx", "EU Tax Observatory, Taxplorer"),
    ("taxplorer-data.json", "https://www.taxplorer.eu/data.json", "EU Tax Observatory, Taxplorer"),
    ("banks-cbcr.xlsx", "https://taxobservatory.world/www-site/uploads/2022/02/Data_cbcr_bank-2.xlsx", "EU Tax Observatory, banks' CbCR database"),
]


def is_public_cbc(dataset):
    text = " ".join([dataset.get("name", ""), dataset.get("title", ""), dataset.get("notes", "")] +
                    [resource.get("name", "") for resource in dataset.get("resources", [])])
    return bool(re.search(r"\bcbc(?:r)?\b|country[\s-]+by[\s-]+country", text, re.I))


def ato_catalogue():
    datasets = []
    urls = []
    while True:
        params = {"fq": f"organization:{ATO_ORGANIZATION}", "rows": 100, "start": len(datasets)}
        url = CATALOGUE_URL + "?" + urllib.parse.urlencode(params)
        with urllib.request.urlopen(url, timeout=40) as response:
            result = json.load(response)["result"]
        urls.append(url)
        datasets.extend(result["results"])
        if len(datasets) >= result["count"]:
            return datasets, urls
        if not result["results"]:
            raise ValueError("ATO catalogue pagination stopped before all datasets were returned")


def inspect_downloads(download_dir):
    with (download_dir / "taxplorer-data.csv").open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        columns = reader.fieldnames
        rows = list(reader)
    if not rows or not {"company", "year", "row_type", "jur_code", "tax_paid"}.issubset(columns):
        raise ValueError("Unexpected Taxplorer CSV schema")
    book = openpyxl.load_workbook(download_dir / "taxplorer-data.xlsx", read_only=True, data_only=True)
    try:
        source = book["Data"].iter_rows(values_only=True)
        if list(next(source)) != columns:
            raise ValueError("Taxplorer CSV and Excel headers differ")
        excel_rows = [["" if value is None else str(value) for value in row] for row in source]
        if excel_rows != [[row[column] for column in columns] for row in rows]:
            raise ValueError("Taxplorer CSV and Excel values differ")
    finally:
        book.close()
    web_data = json.loads((download_dir / "taxplorer-data.json").read_text())
    bank_book = openpyxl.load_workbook(download_dir / "banks-cbcr.xlsx", read_only=True, data_only=True)
    try:
        bank_rows = [row for row in bank_book["DATA"].iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)]
    finally:
        bank_book.close()
    return {
        "taxplorer": {"version": web_data["version"], "csv_excel_rows": len(rows),
                      "companies": len({row["company"] for row in rows}),
                      "company_years": len({(row["company"], row["year"]) for row in rows}),
                      "years": sorted({int(row["year"]) for row in rows}),
                      "row_types": dict(Counter(row["row_type"] for row in rows)),
                      "json_rows": len(web_data["rows"]), "currency": "EUR, full units",
                      "format_note": "CSV and Excel include adjustment rows and full numeric precision. Website JSON omits adjustments and rounds displayed monetary amounts; use CSV/Excel for analysis."},
        "banks": {"records": len(bank_rows), "banks": len({row[3] for row in bank_rows}),
                  "years": sorted({row[2] for row in bank_rows}),
                  "source_labels": dict(Counter(row[-1] for row in bank_rows)),
                  "currency": "EUR millions for financial amounts; staff is a headcount",
                  "units_source": "https://www.taxobservatory.eu/www-site/uploads/2022/12/WP9_Tax-Planning-by-European-Banks_December2022-1.pdf",
                  "units_note": "Original workbook units retained. Its monetary scale differs from Taxplorer's full EUR units."},
    }


def validate(out_dir):
    manifest = json.loads((out_dir / "sources.json").read_text())
    for info in manifest["files"]:
        path = out_dir / info["file"]
        if companies.file_hash(path) != info["sha256"] or path.stat().st_size != info["bytes"]:
            raise ValueError(f"Downloaded file checksum mismatch: {path.name}")
    if inspect_downloads(out_dir / "downloads") != manifest["download_controls"]:
        raise ValueError("CbC download controls differ")
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", type=Path, default=companies.ROOT / "data/cbc")
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    if args.validate_only:
        manifest = validate(args.out_dir)
    else:
        download_dir = args.out_dir / "downloads"
        download_dir.mkdir(parents=True, exist_ok=True)
        datasets, queries = ato_catalogue()
        matches = [dataset for dataset in datasets if is_public_cbc(dataset)]
        files = []
        urls = list(DOWNLOADS)
        for dataset in matches:
            for resource in dataset["resources"]:
                name = resource["id"] + "-" + Path(urllib.parse.urlparse(resource["url"]).path).name
                urls.append((name, resource["url"], "Australian Taxation Office"))
        for name, url, publisher in urls:
            with urllib.request.urlopen(url, timeout=40) as response:
                payload = response.read()
            path = download_dir / name
            path.write_bytes(payload)
            files.append({"file": f"downloads/{name}", "url": url, "publisher": publisher,
                          "bytes": len(payload), "sha256": companies.file_hash(path)})
        manifest = {
            "checked_at_utc": datetime.now(timezone.utc).isoformat(),
            "ato_catalogue": {"queries": queries, "datasets_checked": len(datasets),
                              "matching_datasets": [{"id": dataset["id"], "title": dataset["title"]} for dataset in matches],
                              "guidance_url": GUIDANCE_URL,
                              "interpretation": "No match means no public CbC dataset was found in this catalogue snapshot; it does not establish that no company has published a report elsewhere."},
            "files": files, "download_controls": inspect_downloads(download_dir),
            "taxplorer_source": "https://www.taxplorer.eu/data/",
            "taxplorer_terms": "https://www.taxplorer.eu/licence/",
            "taxplorer_licence": "CC BY-NC-ND 4.0 with publisher-specific terms; republishing the database requires permission and business use requires a licence. Downloads remain local and unchanged.",
            "banks_source": "https://taxobservatory.eu/repository/banks-country-by-country-reporting/",
            "banks_reuse": "No dataset-specific licence identified on the source page. Download retained locally unchanged; not redistributed.",
        }
        (args.out_dir / "sources.json").write_text(json.dumps(manifest, indent=2) + "\n")
        manifest = validate(args.out_dir)
    print(f"ATO public CbC datasets found: {len(manifest['ato_catalogue']['matching_datasets'])}")
    print(f"Taxplorer: {manifest['download_controls']['taxplorer']['csv_excel_rows']:,} unchanged CSV/Excel rows")
    print(f"Banks: {manifest['download_controls']['banks']['records']:,} original workbook records")
    print(f"Downloads verified in {args.out_dir / 'downloads'}")


if __name__ == "__main__":
    main()
