"""Build CSVs of real companies from ATO's public tax-transparency workbook."""
import argparse
import hashlib
import json
import re
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
PACKAGE_URL = "https://data.gov.au/data/api/3/action/package_show?id=corporate-transparency"
DATASET_URL = "https://data.gov.au/data/dataset/corporate-transparency"
LICENCE_URL = "https://creativecommons.org/licenses/by/3.0/au/"
INCOME_HEADERS = ["Name", "ABN", "Total income $", "Taxable income $", "Tax payable $", "Income year"]
PRRT_HEADERS = ["Name", "ABN", "PRRT Payable $"]
INCOME_COLUMNS = ["company_name", "abn", "total_income", "taxable_income", "tax_payable", "income_year"]
PRRT_COLUMNS = ["company_name", "abn", "prrt_payable", "income_year"]
AMOUNT_COLUMNS = ["total_income", "taxable_income", "tax_payable", "prrt_payable"]


def select_resource(package, year=None):
    releases = []
    for resource in package["resources"]:
        match = re.fullmatch(r"(20\d{2}-\d{2}) Report of Entity Tax Information", resource["name"])
        if match:
            releases.append((match[1], resource))
    if year is not None:
        releases = [entry for entry in releases if entry[0] == year]
    if not releases:
        raise ValueError(f"No ATO company release found for {year or 'the latest year'}")
    return max(releases, key=lambda entry: entry[0])


def records_frame(rows, headers, expected_headers, columns):
    if list(headers) != expected_headers:
        raise ValueError(f"Unexpected source columns: {list(headers)!r}")
    frame = pd.DataFrame(rows, columns=columns)
    if frame["company_name"].isna().any() or frame["company_name"].eq("").any():
        raise ValueError("A source record has no company name")
    frame["company_name"] = frame["company_name"].astype("string")
    abns = []
    for value in frame["abn"]:
        if pd.isna(value):
            abns.append(pd.NA)
        else:
            if isinstance(value, str):
                text = value
            else:
                if value != int(value):
                    raise ValueError(f"ABN is not an integer: {value!r}")
                text = str(int(value))
            if not re.fullmatch(r"\d{11}", text):
                raise ValueError(f"ABN is not 11 digits: {text!r}")
            abns.append(text)
    frame["abn"] = pd.array(abns, dtype="string")
    for column in AMOUNT_COLUMNS:
        if column in frame:
            frame[column] = pd.array(frame[column], dtype="Int64")
            if (frame[column].dropna() <= 0).any():
                raise ValueError(f"Unexpected non-positive published amount in {column}")
    if "income_year" in frame:
        frame["income_year"] = frame["income_year"].astype("string")
        if not frame["income_year"].str.fullmatch(r"20\d{2}-\d{2}").fillna(False).all():
            raise ValueError("A source record has an invalid income year")
    return frame


def read_workbook(path, year):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        income_rows = book["Income tax details"].iter_rows(values_only=True)
        headers = next(income_rows)
        income = records_frame(list(income_rows), headers, INCOME_HEADERS, INCOME_COLUMNS)
        prrt_rows = book["PRRT details"].iter_rows(values_only=True)
        headers = next(prrt_rows)
        prrt = records_frame(list(prrt_rows), headers, PRRT_HEADERS, PRRT_COLUMNS[:-1])
        prrt["income_year"] = pd.array([year] * len(prrt), dtype="string")
    finally:
        book.close()
    current = income.loc[income["income_year"].eq(year)].reset_index(drop=True)
    late = income.loc[~income["income_year"].eq(year)].reset_index(drop=True)
    if current.empty:
        raise ValueError(f"Workbook contains no income-tax records for {year}")
    if (late["income_year"] > year).any():
        raise ValueError("Workbook contains a later income year than its release year")
    return {"companies": current, "late_returns": late, "prrt": prrt}


def summary_frame(frames, year):
    income = pd.concat([frames["companies"], frames["late_returns"]], ignore_index=True)
    rows = []
    for income_year, group in income.groupby("income_year", sort=True):
        row = {"release_year": year, "income_year": income_year, "records": len(group),
               "abn_blank_records": int(group["abn"].isna().sum())}
        for column in INCOME_COLUMNS[2:5]:
            row[f"{column}_reported_records"] = int(group[column].count())
            row[f"{column}_reported_sum"] = int(group[column].sum())
            row[f"{column}_blank_records"] = int(group[column].isna().sum())
        rows.append(row)
    return pd.DataFrame(rows)


def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def controls(frame):
    return {
        "records": len(frame),
        "income_years": sorted(frame["income_year"].unique().tolist()),
        "blank_records": {column: int(frame[column].isna().sum()) for column in frame},
        "reported_amount_sums": {column: int(frame[column].sum()) for column in AMOUNT_COLUMNS if column in frame},
    }


def build(workbook, out_dir, year, package=None, resource=None):
    frames = read_workbook(workbook, year)
    suffix = year.replace("-", "_")
    out_dir.mkdir(parents=True, exist_ok=True)
    files = {}
    for kind, frame in frames.items():
        path = out_dir / f"{kind}_{suffix}.csv"
        frame.to_csv(path, index=False)
        files[kind] = {"file": path.name, "sha256": file_hash(path), **controls(frame)}
    summary_path = out_dir / f"summary_{suffix}.csv"
    summary_frame(frames, year).to_csv(summary_path, index=False)
    files["summary"] = {"file": summary_path.name, "sha256": file_hash(summary_path)}
    manifest = {
        "data_kind": "real public corporate tax records, not synthetic",
        "publisher": "Australian Taxation Office, Commonwealth of Australia",
        "release_year": year,
        "dataset_url": DATASET_URL,
        "metadata_url": PACKAGE_URL,
        "licence_url": LICENCE_URL,
        "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_workbook": {"file": workbook.name, "sha256": file_hash(workbook)},
        "amount_units": "annual Australian dollars",
        "blank_amount_meaning": "ATO does not publish amounts of zero or less. Blanks are preserved, not imputed as zero.",
        "coverage": "Corporate income-tax entities meeting ATO's disclosure threshold, plus separately reported PRRT entities. Not all Australian companies or economic groups.",
        "files": files,
    }
    if package is not None:
        manifest["dataset_modified"] = package["metadata_modified"]
        manifest["source_licence"] = package["license_title"]
    if resource is not None:
        manifest["resource"] = {key: resource.get(key) for key in ("id", "name", "url", "last_modified")}
    manifest_path = out_dir / f"source_{suffix}.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest_path


def read_csv(path, columns):
    types = {column: "Int64" if column in AMOUNT_COLUMNS else "string" for column in columns}
    frame = pd.read_csv(
        path, dtype=types, keep_default_na=False,
        na_values={column: [""] for column in columns if column != "company_name"},
    )
    if list(frame.columns) != columns:
        raise ValueError(f"Unexpected CSV columns in {path.name}")
    return frame


def validate(manifest_path, workbook=None):
    manifest = json.loads(manifest_path.read_text())
    year = manifest["release_year"]
    frames = {}
    for kind, columns in (("companies", INCOME_COLUMNS), ("late_returns", INCOME_COLUMNS), ("prrt", PRRT_COLUMNS)):
        expected = manifest["files"][kind]
        path = manifest_path.parent / expected["file"]
        if file_hash(path) != expected["sha256"]:
            raise ValueError(f"CSV checksum mismatch: {path.name}")
        frame = read_csv(path, columns)
        actual = controls(frame)
        if any(actual[key] != expected[key] for key in actual):
            raise ValueError(f"Published source controls do not match: {path.name}")
        frames[kind] = frame
    summary_info = manifest["files"]["summary"]
    summary_path = manifest_path.parent / summary_info["file"]
    if file_hash(summary_path) != summary_info["sha256"]:
        raise ValueError(f"CSV checksum mismatch: {summary_path.name}")
    actual_summary = pd.read_csv(summary_path, dtype={"release_year": "string", "income_year": "string"})
    pd.testing.assert_frame_equal(actual_summary, summary_frame(frames, year), check_dtype=False)
    if workbook is not None:
        if file_hash(workbook) != manifest["source_workbook"]["sha256"]:
            raise ValueError("Source workbook checksum does not match the manifest")
        source = read_workbook(workbook, year)
        for kind in frames:
            pd.testing.assert_frame_equal(frames[kind], source[kind])
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", help="release year, e.g. 2024-25; defaults to latest")
    parser.add_argument("--workbook", type=Path, help="use a downloaded workbook without network access; requires --year")
    parser.add_argument("--raw-dir", type=Path, default=ROOT / "data/raw")
    parser.add_argument("--out-dir", type=Path, default=ROOT / "data/companies")
    parser.add_argument("--validate-only", action="store_true", help="validate existing CSVs offline, optionally against --workbook")
    args = parser.parse_args()
    if args.year is not None and not re.fullmatch(r"20\d{2}-\d{2}", args.year):
        parser.error("--year must use YYYY-YY format")
    if args.workbook is not None and args.year is None:
        parser.error("--workbook requires --year")
    if args.validate_only:
        candidates = sorted(args.out_dir.glob("source_20??_??.json"))
        if args.year:
            manifest_path = args.out_dir / f"source_{args.year.replace('-', '_')}.json"
        else:
            manifest_path = candidates[-1] if candidates else None
        if manifest_path is None or not manifest_path.is_file():
            parser.error("No company data manifest found")
    elif args.workbook is not None:
        manifest_path = build(args.workbook, args.out_dir, args.year)
    else:
        with urllib.request.urlopen(PACKAGE_URL, timeout=60) as response:
            package = json.load(response)["result"]
        year, resource = select_resource(package, args.year)
        args.raw_dir.mkdir(parents=True, exist_ok=True)
        args.workbook = args.raw_dir / Path(resource["url"]).name
        with urllib.request.urlopen(resource["url"], timeout=60) as response:
            args.workbook.write_bytes(response.read())
        manifest_path = build(args.workbook, args.out_dir, year, package, resource)
    manifest = validate(manifest_path, args.workbook)
    print(f"Verified ATO {manifest['release_year']} public company data")
    for kind in ("companies", "late_returns", "prrt"):
        info = manifest["files"][kind]
        print(f"{info['file']}: {info['records']:,} records; published sums {info['reported_amount_sums']}")


if __name__ == "__main__":
    main()
