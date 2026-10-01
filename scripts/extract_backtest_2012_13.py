"""
Extract the genuinely OLD source tables for the "predict at the time" backtest:
the ORIGINAL Taxation Statistics 2012-13 release (not the retrospective 2010-11 to
2022-23 compilation bundled in the current release), plus an older ABS Life Table
vintage. These become the t0 inputs for a 10-year forward simulation whose 2022-23
prediction gets compared against the REAL, already-downloaded 2022-23 ATO figures.

Sources:
  ATO Taxation Statistics 2012-13 (data.gov.au, original per-year release):
    Individuals Table 3  - sex x taxable status x age x income bracket, 67 items
    Individuals Table 14 - percentile distribution of taxable individuals, by sex
  ABS Life Tables, 2016-2018 (abs.gov.au) - national qx by age/sex.
    NOTE: this is the OLDEST vintage retrievable as a direct download (pre-2016
    vintages are archived on the deprecated AUSSTATS site behind JS-driven
    navigation with no scrapeable download links found). Using 2016-2018 for a
    2012-13-start backtest is a minor, deliberate look-ahead compromise -- see
    APPROACH.md. It is still meaningfully older than the 2021-2023 table the main
    model uses.
"""
import re
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed" / "backtest_2012_13"
OUT.mkdir(parents=True, exist_ok=True)


def clean_header(h):
    if h is None:
        return ""
    return re.sub(r"\s+", " ", str(h)).strip()


def extract_table3_2013():
    wb = openpyxl.load_workbook(RAW / "taxstats2013_individual03.xlsx", read_only=True, data_only=True)
    rows = list(wb["Individuals Table 3"].iter_rows(values_only=True))
    header = [clean_header(h) for h in rows[2]]
    dim_cols = header[:4]
    item_cols = header[5:]  # col 4 ("Number of individuals no.") is standalone
    bases = [re.sub(r"\s*no\.?$", "", item_cols[i]).strip() for i in range(0, len(item_cols), 2)]

    records = []
    for r in rows[3:]:
        if r[0] is None:
            continue
        rec = dict(zip(dim_cols, r[:4]))
        rec["Individuals | count"] = r[4]
        for i, base in enumerate(bases):
            c = 5 + 2 * i
            rec[f"{base} | count"] = r[c] if c < len(r) else None
            rec[f"{base} | dollars"] = r[c + 1] if c + 1 < len(r) else None
        records.append(rec)
    df = pd.DataFrame(records)
    df.to_csv(OUT / "table3_2012_13_items_by_bracket.csv", index=False)
    print(f"2012-13 Table 3 -> {len(df)} rows, {len(df.columns)} cols")
    return df


def extract_percentile_2013():
    wb = openpyxl.load_workbook(RAW / "taxstats2013_individual14.xlsx", read_only=True, data_only=True)
    rows = list(wb["Individuals Table 14"].iter_rows(values_only=True))
    gender_row = rows[2]
    item_row = [clean_header(h) for h in rows[3]]

    records = []
    for sex, col_start in [("Female", 1), ("Male", 9)]:
        cols = list(range(col_start, col_start + 8))
        items = item_row[col_start : col_start + 8]
        for pct, r in enumerate(rows[4:104], start=1):
            rec = {"percentile": pct, "sex": sex, "ranged_taxable_income": r[0]}
            for item, c in zip(items, cols):
                rec[item] = r[c] if c < len(r) else None
            records.append(rec)
    df = pd.DataFrame(records)
    df.to_csv(OUT / "table14_2012_13_percentiles.csv", index=False)
    print(f"2012-13 Table 14 (percentiles) -> {len(df)} rows")
    return df


def extract_old_life_table():
    df = pd.read_excel(RAW / "abs_life_tables_2016_2018.xls", sheet_name="Table_1.9", engine="xlrd", header=None)
    # row 6: age header; row 7 onward: data. cols: Age, M_lx, M_qx, M_Lx, M_ex, F_lx, F_qx, F_Lx, F_ex
    data = df.iloc[7:, :9]
    data.columns = ["age", "m_lx", "m_qx", "m_Lx", "m_ex", "f_lx", "f_qx", "f_Lx", "f_ex"]
    data = data[pd.to_numeric(data["age"], errors="coerce").notna()]
    records = []
    for row in data.itertuples():
        records.append({"age": int(row.age), "sex": "Male", "qx": float(row.m_qx)})
        records.append({"age": int(row.age), "sex": "Female", "qx": float(row.f_qx)})
    out = pd.DataFrame(records)
    out.to_csv(OUT / "abs_life_table_2016_2018.csv", index=False)
    print(f"ABS Life Table 2016-2018 (proxy for a pre-2013 vintage) -> {len(out)} rows")


if __name__ == "__main__":
    extract_table3_2013()
    extract_percentile_2013()
    extract_old_life_table()
    print("\nDone. Processed CSVs in", OUT)
